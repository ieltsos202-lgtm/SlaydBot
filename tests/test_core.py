import os
import sys

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config  # noqa: E402
import db  # noqa: E402
import docs  # noqa: E402
import slides  # noqa: E402

PPTX_DATA = {
    "title": "Amir Temur",
    "subtitle": "Davlat boshqaruvi",
    "image_query": "Registan",
    "slides": [
        {"layout": "agenda", "title": "Reja", "bullets": ["Kirish", "Hayoti", "Davlati", "Xulosa"]},
        {"layout": "bullets", "title": "Hayoti", "image_query": "x",
         "bullets": ["Uzun punkt matni " * 3, "Ikkinchi punkt", "Uchinchi punkt"]},
        {"layout": "cards", "title": "Omillar", "items": [{"title": "A", "text": "aa"}, {"title": "B", "text": "bb"}]},
        {"layout": "stats", "title": "Raqamlar", "stats": [{"value": "1370", "label": "yil"}, {"value": "35", "label": "yil"}],
         "bullets": ["Izoh"]},
        {"layout": "timeline", "title": "Tarix", "events": [{"year": str(y), "text": "voqea"} for y in (1336, 1370, 1405)]},
        {"layout": "quote", "title": "Iqtibos", "image_query": "x", "quote": {"text": "Kuch adolatdadir", "author": "Amir Temur"}},
        {"layout": "two_column", "title": "Taqqoslash", "left": {"title": "A", "bullets": ["a"]},
         "right": {"title": "B", "bullets": ["b"]}},
        {"layout": "stats", "title": "Buzuq", "stats": []},
    ],
}
DOCX_DATA = {
    "title": "Amir Temur",
    "intro": ["Kirish abzatsi."],
    "sections": [{"heading": "Bo'lim", "paragraphs": ["Matn."]}],
    "conclusion": ["Xulosa."],
    "references": ["Manba 1"],
}


@pytest.mark.parametrize("theme", list(slides.THEMES))
def test_build_pptx_all_layouts(tmp_path, theme):
    img = tmp_path / "i.jpg"
    Image.new("RGB", (1600, 900), "steelblue").save(img)
    pics = {"title": str(img), 1: str(img), 5: str(img)}
    path = tmp_path / "a.pptx"
    slides.build_pptx(PPTX_DATA, theme, "uz", str(path), pics)
    assert len(Presentation(str(path)).slides) == 10


def test_build_pptx_without_images(tmp_path):
    path = tmp_path / "a.pptx"
    slides.build_pptx(PPTX_DATA, "ocean", "ru", str(path))
    assert len(Presentation(str(path)).slides) == 10


def test_normalize_fallback():
    assert slides.normalize({"layout": "stats", "title": "X", "stats": []})["layout"] == "bullets"
    assert slides.normalize({"layout": "weird", "bullets": ["a", "b"]})["layout"] == "bullets"


def test_image_queries():
    q = slides.image_queries(PPTX_DATA, "topic")
    assert q == {"title": "Registan", 1: "x", 5: "x"}


def test_wiki_relevance():
    import wiki
    assert wiki.relevance("Amir Temur va uning davlati", "Amir Temur") >= 0.5
    assert wiki.relevance("Zamonaviy texnologiyalar va kiberxavfsizlik",
                          "Toshkent axborot texnologiyalari universiteti") < 0.5


def test_build_docx(tmp_path):
    path = tmp_path / "a.docx"
    docs.build_docx(DOCX_DATA, "ru", str(path))
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "ВВЕДЕНИЕ" in text and "1. Bo'lim" in text


def test_parse_amount():
    import gemini
    assert gemini.parse_amount(5000) == 5000
    assert gemini.parse_amount("5 000,00 so'm") == 5000
    assert gemini.parse_amount("12,000.00") == 12000
    assert gemini.parse_amount("35.000") == 35000
    assert gemini.parse_amount(None) is None
    assert gemini.card_matches({"card_last4": "9860 16** **** 3004"}, "9860160144093004")


def test_package_for_amount(monkeypatch):
    import handlers_user
    monkeypatch.setattr(config, "PACKAGES", {1: 5000, 3: 12000, 10: 35000})
    assert handlers_user.package_for_amount(12000) == (3, 12000)
    assert handlers_user.package_for_amount(20000) == (3, 12000)
    assert handlers_user.package_for_amount(4000) is None
    assert handlers_user.package_for_amount(None) is None


def test_safe_filename():
    assert docs.safe_filename('a/b:c*"?', "pptx") == "abc.pptx"
    assert docs.safe_filename("///", "docx") == "hujjat.docx"


@pytest.fixture
async def fresh_db(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DB_PATH", str(tmp_path / "t.db"))
    monkeypatch.setattr(config, "FREE_CREDITS", 1)
    monkeypatch.setattr(config, "REFERRAL_BONUS", 1)
    await db.init()


async def test_credits_and_referral(fresh_db):
    await db.get_or_create_user(1, "a", "A")
    user, is_new = await db.get_or_create_user(2, "b", "B", referred_by=1)
    assert is_new and user["referred_by"] == 1 and user["credits"] == 1

    assert await db.spend_credit(2)
    assert not await db.spend_credit(2)

    pid = await db.create_payment(2, 5000, 1, "f", "u1", "tx1", "")
    payment, referrer = await db.approve_payment(pid)
    assert payment and referrer == 1
    assert (await db.get_user(2))["credits"] == 1
    assert (await db.get_user(1))["credits"] == 2

    again, _ = await db.approve_payment(pid)
    assert again is None
    assert await db.receipt_used("u1", None)
    assert await db.receipt_used("other", "tx1")

    pid2 = await db.create_payment(2, 5000, 1, "f", "u2", None, "")
    _, referrer2 = await db.approve_payment(pid2)
    assert referrer2 is None
    assert (await db.stats())["revenue"] == 10000


async def test_self_referral_ignored(fresh_db):
    user, _ = await db.get_or_create_user(5, None, "X", referred_by=5)
    assert user["referred_by"] is None
