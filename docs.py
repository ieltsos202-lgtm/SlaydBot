import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "1F3864"
ACCENT = "2E75B6"
FONT = "Times New Roman"

LABELS = {
    "uz": {"ministry": "O'ZBEKISTON RESPUBLIKASI\nOLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI",
           "uni": "____________________ UNIVERSITETI", "faculty": "«__________» fakulteti",
           "plan": "REJA", "intro": "KIRISH", "conclusion": "XULOSA", "refs": "FOYDALANILGAN ADABIYOTLAR",
           "type": "REFERAT", "topic": "Mavzu", "done": "Bajardi", "group": "Guruh",
           "checked": "Qabul qildi", "city": "Toshkent"},
    "ru": {"ministry": "МИНИСТЕРСТВО ВЫСШЕГО ОБРАЗОВАНИЯ, НАУКИ И ИННОВАЦИЙ\nРЕСПУБЛИКИ УЗБЕКИСТАН",
           "uni": "____________________ УНИВЕРСИТЕТ", "faculty": "Факультет «__________»",
           "plan": "СОДЕРЖАНИЕ", "intro": "ВВЕДЕНИЕ", "conclusion": "ЗАКЛЮЧЕНИЕ", "refs": "СПИСОК ЛИТЕРАТУРЫ",
           "type": "РЕФЕРАТ", "topic": "Тема", "done": "Выполнил(а)", "group": "Группа",
           "checked": "Принял(а)", "city": "Ташкент"},
    "en": {"ministry": "MINISTRY OF HIGHER EDUCATION, SCIENCE AND INNOVATION\nOF THE REPUBLIC OF UZBEKISTAN",
           "uni": "____________________ UNIVERSITY", "faculty": "Faculty of __________",
           "plan": "CONTENTS", "intro": "INTRODUCTION", "conclusion": "CONCLUSION", "refs": "REFERENCES",
           "type": "ESSAY", "topic": "Topic", "done": "Prepared by", "group": "Group",
           "checked": "Checked by", "city": "Tashkent"},
}

_PPR_AFTER_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                  "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                  "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                  "w:pPrChange")


def safe_filename(topic: str, ext: str) -> str:
    name = re.sub(r"[^\w\- ]", "", topic, flags=re.U).strip()[:50] or "hujjat"
    return f"{name}.{ext}"


def _para(doc, text: str, size=14, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          indent=True, space_after=6, color=None, line_spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    fmt = p.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(0)
    if indent:
        fmt.first_line_indent = Cm(1.25)
    for i, line in enumerate(str(text).split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = FONT
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def _bottom_border(p, color=ACCENT, size=12):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "4"), ("w:color", color)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.insert_element_before(bdr, *_PPR_AFTER_BDR)


def _heading(doc, text: str, page_break=False):
    if page_break:
        _page_break(doc)
    p = _para(doc, text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
              space_after=14, color=NAVY, line_spacing=1.2)
    p.paragraph_format.keep_with_next = True
    _bottom_border(p)
    return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _page_border(section):
    sectPr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    borders.set(qn("w:display"), "firstPage")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        for k, v in (("w:val", "thinThickSmallGap"), ("w:sz", "24"), ("w:space", "24"), ("w:color", NAVY)):
            el.set(qn(k), v)
        borders.append(el)
    sectPr.find(qn("w:pgMar")).addnext(borders)


def _page_number(paragraph):
    run = paragraph.add_run()
    run.font.name = FONT
    run.font.size = Pt(12)
    for kind, instr in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)


def build_docx(data: dict, lang: str, path: str) -> None:
    lb = LABELS.get(lang, LABELS["uz"])
    title = str(data.get("title", "")).strip()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.different_first_page_header_footer = True
    _page_border(sec)
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_number(footer_p)
    header_p = sec.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header_p.add_run(title[:90])
    hr.font.name, hr.font.size, hr.font.italic = FONT, Pt(10), True
    hr.font.color.rgb = RGBColor.from_string("7F7F7F")

    center = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, lb["ministry"], size=13, bold=True, align=center, indent=False, color=NAVY, line_spacing=1.15)
    _para(doc, lb["uni"], size=13, bold=True, align=center, indent=False, line_spacing=1.15)
    _para(doc, lb["faculty"], size=13, align=center, indent=False, line_spacing=1.15)
    for _ in range(5):
        doc.add_paragraph()
    _para(doc, lb["type"], size=36, bold=True, align=center, indent=False, color=NAVY, line_spacing=1.0)
    deco = _para(doc, "", size=6, align=center, indent=False, space_after=12)
    _bottom_border(deco, ACCENT, 18)
    _para(doc, f"{lb['topic']}:", size=14, italic=True, align=center, indent=False, line_spacing=1.0)
    _para(doc, f"«{title}»", size=20, bold=True, align=center, indent=False, color=ACCENT, line_spacing=1.2)
    for _ in range(5):
        doc.add_paragraph()
    right = WD_ALIGN_PARAGRAPH.RIGHT
    for key in ("done", "group", "checked"):
        _para(doc, f"{lb[key]}: ______________________", size=14, align=right, indent=False, line_spacing=1.3)
    for _ in range(3):
        doc.add_paragraph()
    _para(doc, f"{lb['city']} — {date.today().year}", size=14, bold=True, align=center, indent=False)

    sections = [s for s in data.get("sections", []) if isinstance(s, dict)]
    _heading(doc, lb["plan"], page_break=True)
    plan = [lb["intro"]] + [f"{i}. {s.get('heading', '')}" for i, s in enumerate(sections, 1)]
    plan += [lb["conclusion"], lb["refs"]]
    for item in plan:
        p = _para(doc, item, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_after=4)
        p.paragraph_format.left_indent = Cm(1)

    _heading(doc, lb["intro"], page_break=True)
    for text in data.get("intro", []):
        _para(doc, str(text))

    for i, s in enumerate(sections, 1):
        _heading(doc, f"{i}. {s.get('heading', '')}", page_break=True)
        for text in s.get("paragraphs", []):
            _para(doc, str(text))

    _heading(doc, lb["conclusion"], page_break=True)
    for text in data.get("conclusion", []):
        _para(doc, str(text))

    _heading(doc, lb["refs"], page_break=True)
    for i, ref in enumerate(data.get("references", []), 1):
        p = _para(doc, f"{i}. {ref}", align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)

    doc.save(path)
