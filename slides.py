import math

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W, H = 13.333, 7.5
RECT, ROUND, OVAL = MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL
LEFT, CENTER, RIGHT = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP, MIDDLE, BOTTOM = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM

_BASE = {"head": "Segoe UI Semibold", "body": "Segoe UI", "on_dark": "FFFFFF", "card": "FFFFFF"}
THEMES = {
    "ocean": {**_BASE, "name": "🌊 Okean", "dark1": "0B1E3F", "dark2": "1456A0", "accent": "00B4F0",
              "accent2": "36D6B5", "bg": "F4F8FD", "title": "0B1E3F", "text": "1B2A41",
              "muted": "55657D", "soft": "BFD8F5", "border": "DCE6F2"},
    "sunset": {**_BASE, "name": "🌅 Quyosh", "dark1": "2B1055", "dark2": "B23A6F", "accent": "FF7A45",
               "accent2": "FFB547", "bg": "FFF8F4", "title": "2B1055", "text": "2E2440",
               "muted": "6E6280", "soft": "F6CFE0", "border": "F1DFD6"},
    "emerald": {**_BASE, "name": "🌿 Zumrad", "dark1": "042F2A", "dark2": "0E7A6B", "accent": "10B981",
                "accent2": "F5B700", "bg": "F3FAF7", "title": "053B33", "text": "1D2E2A",
                "muted": "56706A", "soft": "BDEBDD", "border": "D5EAE2"},
    "gold": {**_BASE, "name": "👑 Oltin", "head": "Georgia", "dark1": "08090C", "dark2": "262B36",
             "accent": "D4AF37", "accent2": "F1D98B", "bg": "111318", "card": "1B1F27",
             "title": "F1D98B", "text": "ECECEC", "muted": "A4A9B4", "soft": "D9D2B8", "border": "2C313C"},
    "royal": {**_BASE, "name": "💜 Qirollik", "dark1": "2A0845", "dark2": "6A1B9A", "accent": "E94BA0",
              "accent2": "9F7AEA", "bg": "FAF6FF", "title": "2A0845", "text": "2B2140",
              "muted": "6B5E80", "soft": "E2CCF5", "border": "E9DDF5"},
    "minimal": {**_BASE, "name": "⚪ Minimal", "dark1": "121212", "dark2": "3A3A3A", "accent": "E63946",
                "accent2": "1D7FB8", "bg": "FFFFFF", "card": "F5F6F8", "title": "121212",
                "text": "222222", "muted": "666B73", "soft": "D6D6D6", "border": "E4E6EA"},
}

LABELS = {
    "uz": {"kicker": "TAQDIMOT", "thanks": "E'tiboringiz uchun rahmat!",
           "questions": "Savollaringiz bo'lsa, marhamat"},
    "ru": {"kicker": "ПРЕЗЕНТАЦИЯ", "thanks": "Спасибо за внимание!",
           "questions": "Буду рад ответить на ваши вопросы"},
    "en": {"kicker": "PRESENTATION", "thanks": "Thank you for your attention!",
           "questions": "Your questions are welcome"},
}

IMAGE_LAYOUTS = {"bullets", "quote"}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _alpha(fill, opacity: float) -> None:
    clr = fill._xPr.find(qn("a:solidFill"))[0]
    el = etree.SubElement(clr, qn("a:alpha"))
    el.set("val", str(int(opacity * 100000)))


def _grad(fill, c1: str, c2: str, angle: float = 45) -> None:
    fill.gradient()
    fill.gradient_angle = angle
    stops = fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = _rgb(c1)
    stops[-1].position = 1.0
    stops[-1].color.rgb = _rgb(c2)


def _bg(slide, color: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _bg_grad(slide, c1: str, c2: str, angle: float = 45) -> None:
    _grad(slide.background.fill, c1, c2, angle)


def _shape(slide, kind, x, y, w, h, color=None, alpha=None, line=None, radius=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if color:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(color)
        if alpha is not None:
            _alpha(shp.fill, alpha)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if radius is not None and kind == ROUND:
        shp.adjustments[0] = radius
    return shp


def _fit(texts: list[str], w: float, h: float, max_size: int, min_size: int,
         spacing: float = 1.2, gap_pt: float = 0) -> int:
    for size in range(max_size, min_size - 1, -1):
        cpl = max(1, int(w * 72 / (size * 0.52)))
        lines = sum(max(1, math.ceil(len(t) / cpl)) for t in texts)
        need = lines * size * spacing / 72 + gap_pt * max(0, len(texts) - 1) / 72
        if need <= h:
            return size
    return min_size


def _text(slide, x, y, w, h, text: str, size: int, color: str, font: str, bold=False,
          italic=False, align=LEFT, anchor=TOP, min_size=None, spacing=1.1, tracking=None):
    lines = str(text).split("\n")
    if min_size:
        size = _fit(lines, w, h, size, min_size, spacing + 0.1)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.name = font
        f.color.rgb = _rgb(color)
        if tracking:
            run._r.get_or_add_rPr().set("spc", str(tracking))
    return box


def _bullet_box(slide, x, y, w, h, items: list[str], th: dict, marker: str, max_size=20, min_size=12):
    gap = 10
    size = _fit(items, w - 0.35, h, max_size, min_size, 1.3, gap)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        p.space_after = Pt(gap)
        m = p.add_run()
        m.text = f"{marker}  "
        m.font.size = Pt(size)
        m.font.bold = True
        m.font.color.rgb = _rgb(th["accent"])
        r = p.add_run()
        r.text = item
        r.font.size = Pt(size)
        r.font.name = th["body"]
        r.font.color.rgb = _rgb(th["text"])
    return box


def _picture(slide, path: str, x, y, w, h, rounded=False):
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        c = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = c
    if rounded:
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        if geom is not None:
            geom.set("prst", "roundRect")
    return pic


def _strs(value, limit: int) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()][:limit]


def _pairs(value, k1: str, k2: str, limit: int) -> list[dict]:
    if not isinstance(value, list):
        return []
    out = []
    for item in value:
        if isinstance(item, dict) and str(item.get(k1, "")).strip() and str(item.get(k2, "")).strip():
            out.append({k1: str(item[k1]).strip(), k2: str(item[k2]).strip()})
    return out[:limit]


def normalize(sl: dict) -> dict:
    layout = str(sl.get("layout", "bullets")).strip().lower()
    title = str(sl.get("title", "")).strip()
    bullets = _strs(sl.get("bullets"), 6)
    out = {"layout": layout, "title": title, "bullets": bullets}
    if layout == "cards":
        out["items"] = _pairs(sl.get("items"), "title", "text", 4)
        ok = len(out["items"]) >= 2
    elif layout == "stats":
        out["stats"] = _pairs(sl.get("stats"), "value", "label", 4)
        ok = len(out["stats"]) >= 2
    elif layout == "timeline":
        out["events"] = _pairs(sl.get("events"), "year", "text", 6)
        ok = len(out["events"]) >= 3
    elif layout == "quote":
        q = sl.get("quote") if isinstance(sl.get("quote"), dict) else {}
        out["quote"] = {"text": str(q.get("text", "")).strip(), "author": str(q.get("author", "")).strip()}
        ok = bool(out["quote"]["text"])
    elif layout == "two_column":
        cols = []
        for key in ("left", "right"):
            c = sl.get(key) if isinstance(sl.get(key), dict) else {}
            cols.append({"title": str(c.get("title", "")).strip(), "bullets": _strs(c.get("bullets"), 5)})
        out["left"], out["right"] = cols
        ok = all(c["title"] and c["bullets"] for c in cols)
    elif layout in ("agenda", "bullets"):
        ok = len(bullets) >= 2
    else:
        ok = False
    if not ok:
        out["layout"] = "bullets"
        if not out["bullets"]:
            extra = [f"{i.get('title', '')}: {i.get('text', '')}" for i in out.get("items", [])]
            extra += [f"{s.get('value', '')} — {s.get('label', '')}" for s in out.get("stats", [])]
            extra += [f"{e.get('year', '')}: {e.get('text', '')}" for e in out.get("events", [])]
            out["bullets"] = extra[:6] or [title or "—"]
    return out


def _decor_dark(slide, th: dict) -> None:
    _shape(slide, OVAL, W - 4.2, -2.4, 6.5, 6.5, th["accent"], alpha=0.10)
    _shape(slide, OVAL, W - 2.2, 4.6, 4.2, 4.2, th["accent2"], alpha=0.08)
    _shape(slide, OVAL, -1.6, H - 2.2, 3.6, 3.6, th["accent"], alpha=0.12)


def _content_base(slide, th: dict, title: str, n: int, total: int) -> None:
    _bg(slide, th["bg"])
    _shape(slide, OVAL, W - 3.0, -2.2, 5.0, 5.0, th["accent"], alpha=0.07)
    _shape(slide, OVAL, W - 1.5, -0.8, 2.4, 2.4, th["accent2"], alpha=0.10)
    _shape(slide, ROUND, 0.6, 0.6, 0.11, 0.8, th["accent"], radius=0.5)
    _text(slide, 0.9, 0.45, 10.6, 1.1, title, 34, th["title"], th["head"], bold=True,
          anchor=MIDDLE, min_size=22)
    _shape(slide, RECT, 0.6, H - 0.42, 1.1, 0.05, th["accent"])
    _shape(slide, RECT, 1.75, H - 0.42, 0.35, 0.05, th["accent2"])
    _text(slide, W - 1.9, H - 0.62, 1.3, 0.35, f"{n:02d} / {total:02d}", 11, th["muted"], th["body"],
          align=RIGHT, tracking=100)


def _title_slide(slide, th, lb, title, subtitle, img):
    _bg_grad(slide, th["dark1"], th["dark2"], 30)
    if img:
        _picture(slide, img, 7.4, 0, W - 7.4, H)
        _shape(slide, RECT, 7.4, 0, W - 7.4, H, th["dark1"], alpha=0.18)
        _shape(slide, RECT, 7.25, 0, 0.15, H, th["accent"])
        tw = 6.0
        _shape(slide, OVAL, -1.4, H - 2.6, 3.6, 3.6, th["accent"], alpha=0.14)
        _shape(slide, OVAL, 5.6, -1.0, 2.4, 2.4, th["accent2"], alpha=0.10)
    else:
        _decor_dark(slide, th)
        tw = 9.6
    _shape(slide, ROUND, 0.9, 1.55, 0.5, 0.08, th["accent"], radius=0.5)
    _text(slide, 1.5, 1.4, 5.0, 0.4, lb["kicker"], 13, th["accent"], th["body"], bold=True, tracking=500)
    _text(slide, 0.9, 1.95, tw, 2.9, title, 50, th["on_dark"], th["head"], bold=True,
          anchor=BOTTOM, min_size=28, spacing=1.0)
    _shape(slide, RECT, 0.9, 5.05, 1.6, 0.08, th["accent"])
    _shape(slide, RECT, 2.55, 5.05, 0.4, 0.08, th["accent2"])
    if subtitle:
        _text(slide, 0.9, 5.35, tw, 1.3, subtitle, 20, th["soft"], th["body"], min_size=13)


def _grid(slide, th, items: list[str]):
    n = len(items)
    cols = 2 if n > 3 else 1
    rows = math.ceil(n / cols)
    x0, y0, aw, ah, gap = 0.9, 1.85, 11.6, 5.0, 0.28
    cw = (aw - gap * (cols - 1)) / cols
    ch = min(2.45, (ah - gap * (rows - 1)) / rows)
    y0 += (ah - (ch * rows + gap * (rows - 1))) / 2
    size = min(_fit([t], cw - 1.75, ch - 0.4, 24, 13, 1.25) for t in items)
    for i, t in enumerate(items):
        r, c = divmod(i, cols)
        x, y = x0 + c * (cw + gap), y0 + r * (ch + gap)
        _shape(slide, ROUND, x, y, cw, ch, th["card"], line=th["border"], radius=0.09)
        _shape(slide, ROUND, x, y + 0.2, 0.09, ch - 0.4, th["accent"] if i % 2 == 0 else th["accent2"], radius=0.5)
        d = min(0.78, ch - 0.4)
        _shape(slide, OVAL, x + 0.3, y + (ch - d) / 2, d, d, th["accent"] if i % 2 == 0 else th["accent2"])
        _text(slide, x + 0.3, y + (ch - d) / 2, d, d, f"{i + 1:02d}", 17, th["on_dark"], th["head"],
              bold=True, align=CENTER, anchor=MIDDLE)
        _text(slide, x + 0.5 + d, y + 0.12, cw - d - 0.75, ch - 0.24, t, size, th["text"], th["body"],
              anchor=MIDDLE)


def _bullets(slide, th, bullets: list[str], img, flip: bool):
    if not img:
        _grid(slide, th, bullets)
        return
    if flip:
        ix, tx = 0.9, 6.25
    else:
        ix, tx = 8.0, 0.9
    tw, iw, ih, iy = 6.2, 4.45, 4.85, 1.8
    _shape(slide, ROUND, ix + 0.28, iy + 0.28, iw, ih, th["accent"], alpha=0.30, radius=0.08)
    _picture(slide, img, ix, iy, iw, ih, rounded=True)
    n = len(bullets)
    gap, ah = 0.16, 4.95
    row_h = (ah - gap * (n - 1)) / n
    size = min(_fit([b], tw - 0.85, row_h - 0.05, 20, 11, 1.25) for b in bullets)
    for i, b in enumerate(bullets):
        y = 1.8 + i * (row_h + gap)
        d = min(0.5, row_h * 0.7)
        color = th["accent"] if i % 2 == 0 else th["accent2"]
        _shape(slide, OVAL, tx, y + (row_h - d) / 2, d, d, color)
        _text(slide, tx, y + (row_h - d) / 2, d, d, str(i + 1), 14, th["on_dark"], th["head"],
              bold=True, align=CENTER, anchor=MIDDLE)
        _text(slide, tx + d + 0.25, y, tw - d - 0.3, row_h, b, size, th["text"], th["body"], anchor=MIDDLE)


def _cards(slide, th, items: list[dict]):
    n = len(items)
    x0, y0, aw, ah, gap = 0.9, 1.9, 11.6, 4.95, 0.3
    cw = (aw - gap * (n - 1)) / n
    tsize = min(_fit([it["title"]], cw - 0.6, 0.95, 22, 13, 1.2) for it in items)
    bsize = min(_fit([it["text"]], cw - 0.7, ah - 2.85, 18, 10, 1.3) for it in items)
    for i, it in enumerate(items):
        x = x0 + i * (cw + gap)
        color = th["accent"] if i % 2 == 0 else th["accent2"]
        _shape(slide, ROUND, x, y0, cw, ah, th["card"], line=th["border"], radius=0.07)
        _shape(slide, ROUND, x + 0.35, y0 + 0.4, 0.9, 0.09, color, radius=0.5)
        _text(slide, x + 0.35, y0 + 0.6, cw - 0.7, 0.85, f"{i + 1:02d}", 38, color, th["head"], bold=True)
        _text(slide, x + 0.35, y0 + 1.5, cw - 0.7, 0.95, it["title"], tsize, th["title"], th["head"],
              bold=True, anchor=BOTTOM)
        _text(slide, x + 0.35, y0 + 2.6, cw - 0.7, ah - 2.85, it["text"], bsize, th["muted"], th["body"],
              spacing=1.2)


def _stats(slide, th, stats: list[dict], bullets: list[str]):
    n = len(stats)
    x0, aw, gap = 0.9, 11.6, 0.3
    y, ch = (2.0, 3.3) if bullets else (2.25, 3.9)
    cw = (aw - gap * (n - 1)) / n
    longest = max(len(st["value"]) for st in stats)
    vsize = max(20, min(54, int((cw - 0.7) * 72 / (longest * 0.62))))
    for i, st in enumerate(stats):
        x = x0 + i * (cw + gap)
        card = _shape(slide, ROUND, x, y, cw, ch, radius=0.08)
        _grad(card.fill, th["dark1"], th["dark2"], 315 if i % 2 else 45)
        _shape(slide, OVAL, x + cw - 1.1, y + 0.25, 0.8, 0.8, th["accent"], alpha=0.25)
        _shape(slide, RECT, x + 0.4, y + 0.45, 0.7, 0.07, th["accent"])
        _text(slide, x + 0.35, y + 0.7, cw - 0.7, 1.45, st["value"], vsize, th["accent2"], th["head"],
              bold=True, anchor=MIDDLE, spacing=1.0)
        _text(slide, x + 0.4, y + 2.2, cw - 0.8, ch - 2.4, st["label"], 18, th["soft"], th["body"],
              min_size=11)
    if bullets:
        _text(slide, 0.9, y + ch + 0.35, 11.6, 6.85 - (y + ch + 0.35), " • ".join(bullets[:2]), 16,
              th["muted"], th["body"], align=CENTER, anchor=MIDDLE, min_size=11)


def _timeline(slide, th, events: list[dict]):
    n = len(events)
    x0, x1, ly = 0.9, 12.45, 3.55
    seg = (x1 - x0) / n
    _shape(slide, ROUND, x0, ly - 0.04, x1 - x0, 0.08, th["border"], radius=0.5)
    _shape(slide, ROUND, x0, ly - 0.04, (x1 - x0) * 0.5, 0.08, th["accent"], alpha=0.6, radius=0.5)
    tsize = min(_fit([e["text"]], seg - 0.56, 2.35, 18, 10, 1.3) for e in events)
    for i, ev in enumerate(events):
        cx = x0 + seg * i + seg / 2
        color = th["accent"] if i % 2 == 0 else th["accent2"]
        _text(slide, cx - seg / 2 + 0.08, ly - 1.35, seg - 0.16, 1.0, ev["year"], 26, color, th["head"],
              bold=True, align=CENTER, anchor=BOTTOM, min_size=13)
        _shape(slide, OVAL, cx - 0.24, ly - 0.24, 0.48, 0.48, color)
        _shape(slide, OVAL, cx - 0.1, ly - 0.1, 0.2, 0.2, th["bg"])
        _shape(slide, RECT, cx - 0.015, ly + 0.25, 0.03, 0.35, color)
        _shape(slide, ROUND, cx - seg / 2 + 0.1, ly + 0.6, seg - 0.2, 2.65, th["card"], line=th["border"],
               radius=0.1)
        _text(slide, cx - seg / 2 + 0.28, ly + 0.75, seg - 0.56, 2.35, ev["text"], tsize, th["text"],
              th["body"], align=CENTER, anchor=MIDDLE)


def _two_column(slide, th, left: dict, right: dict):
    w, gap, y, h = 5.65, 0.3, 1.85, 5.0
    for i, col in enumerate((left, right)):
        x = 0.9 + i * (w + gap)
        _shape(slide, ROUND, x, y, w, h, th["card"], line=th["border"], radius=0.06)
        head = _shape(slide, ROUND, x, y, w, 1.0, radius=0.2)
        _grad(head.fill, th["dark1"], th["dark2"] if i == 0 else th["accent2"], 0)
        _text(slide, x + 0.35, y, w - 0.7, 1.0, col["title"], 22, th["on_dark"], th["head"], bold=True,
              anchor=MIDDLE, min_size=13)
        _bullet_box(slide, x + 0.35, y + 1.3, w - 0.7, h - 1.55, col["bullets"], th, "✔", 19, 11)


def _quote(slide, th, title: str, quote: dict, img):
    _bg_grad(slide, th["dark1"], th["dark2"], 30)
    if img:
        _picture(slide, img, 0, 0, 4.9, H)
        _shape(slide, RECT, 4.9, 0, 0.12, H, th["accent"])
        tx, tw = 5.75, 6.9
    else:
        _decor_dark(slide, th)
        tx, tw = 1.4, 10.4
    if title:
        _text(slide, tx, 0.75, tw, 0.5, title.upper(), 13, th["soft"], th["body"], bold=True, tracking=300)
    _text(slide, tx - 0.1, 1.05, 2.0, 1.6, "\u201C", 150, th["accent"], "Georgia", bold=True)
    _text(slide, tx, 2.5, tw, 3.2, quote["text"], 30, th["on_dark"], "Georgia", italic=True,
          anchor=MIDDLE, min_size=16, spacing=1.15)
    if quote["author"]:
        _shape(slide, RECT, tx, 6.0, 0.9, 0.06, th["accent"])
        _text(slide, tx, 6.15, tw, 0.6, "— " + quote["author"], 18, th["accent2"], th["body"], bold=True)


def _thanks(slide, th, lb, img):
    _bg_grad(slide, th["dark1"], th["dark2"], 30)
    if img:
        _picture(slide, img, 0, 0, W, H)
        _shape(slide, RECT, 0, 0, W, H, th["dark1"], alpha=0.78)
    _decor_dark(slide, th)
    _text(slide, 1.0, 2.2, W - 2.0, 1.7, lb["thanks"], 54, th["on_dark"], th["head"], bold=True,
          align=CENTER, anchor=MIDDLE, min_size=30)
    _shape(slide, RECT, W / 2 - 0.9, 4.1, 1.3, 0.08, th["accent"])
    _shape(slide, RECT, W / 2 + 0.45, 4.1, 0.45, 0.08, th["accent2"])
    _text(slide, 1.0, 4.4, W - 2.0, 0.8, lb["questions"], 22, th["soft"], th["body"], align=CENTER)


def build_pptx(data: dict, theme_key: str, lang: str, path: str, images: dict | None = None) -> None:
    images = dict(images or {})
    if not images.get("title"):
        spare = [p for k, p in images.items() if k != "title" and p]
        if spare:
            images["title"] = spare[-1]
    th = THEMES.get(theme_key, THEMES["ocean"])
    lb = LABELS.get(lang, LABELS["uz"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]

    _title_slide(prs.slides.add_slide(blank), th, lb, str(data.get("title", "")).strip(),
                 str(data.get("subtitle", "")).strip(), images.get("title"))

    slides = [normalize(s) for s in data.get("slides", []) if isinstance(s, dict)]
    total = len(slides)
    image_count = 0
    for i, sl in enumerate(slides):
        s = prs.slides.add_slide(blank)
        img = images.get(i) if sl["layout"] in IMAGE_LAYOUTS else None
        if sl["layout"] == "quote":
            _quote(s, th, sl["title"], sl["quote"], img)
            continue
        _content_base(s, th, sl["title"], i + 1, total)
        layout = sl["layout"]
        if layout == "agenda":
            _grid(s, th, sl["bullets"])
        elif layout == "cards":
            _cards(s, th, sl["items"])
        elif layout == "stats":
            _stats(s, th, sl["stats"], sl["bullets"])
        elif layout == "timeline":
            _timeline(s, th, sl["events"])
        elif layout == "two_column":
            _two_column(s, th, sl["left"], sl["right"])
        else:
            _bullets(s, th, sl["bullets"], img, flip=image_count % 2 == 1)
            if img:
                image_count += 1

    _thanks(prs.slides.add_slide(blank), th, lb, images.get("title"))
    prs.save(path)


def image_queries(data: dict, topic: str) -> dict:
    queries = {"title": str(data.get("image_query") or topic)}
    for i, sl in enumerate(data.get("slides", [])):
        if isinstance(sl, dict) and sl.get("image_query") and normalize(sl)["layout"] in IMAGE_LAYOUTS:
            queries[i] = str(sl["image_query"])
    return queries
